import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "customers.xlsx"

# ---------------- Excel File ----------------
def create_file():
    if not os.path.exists(FILE_NAME):
        wb = Workbook()
        ws = wb.active
        ws.title = "Customers"
        ws.append(["Customer ID", "Name", "Phone", "Email", "Address"])
        wb.save(FILE_NAME)


# ---------------- Login Window ----------------
def login():
    username = username_entry.get()
    password = password_entry.get()

    if username == "admin" and password == "admin123":
        login_window.destroy()
        customer_system()
    else:
        messagebox.showerror("Login Error", "Invalid username or password")


login_window = tk.Tk()
login_window.title("Customer Management System - Login")
login_window.geometry("400x300")

tk.Label(
    login_window,
    text="CUSTOMER MANAGEMENT SYSTEM",
    font=("Arial", 16, "bold")
).pack(pady=25)

tk.Label(login_window, text="Username").pack()
username_entry = tk.Entry(login_window, width=30)
username_entry.pack(pady=5)

tk.Label(login_window, text="Password").pack()
password_entry = tk.Entry(login_window, width=30, show="*")
password_entry.pack(pady=5)

tk.Button(
    login_window,
    text="Login",
    width=15,
    command=login
).pack(pady=20)


# ---------------- Customer System ----------------
def customer_system():

    root = tk.Tk()
    root.title("Customer Management System")
    root.geometry("950x600")

    # Variables
    customer_id = tk.StringVar()
    name = tk.StringVar()
    phone = tk.StringVar()
    email = tk.StringVar()
    address = tk.StringVar()
    search_var = tk.StringVar()

    # -------- Functions --------

    def clear_fields():
        customer_id.set("")
        name.set("")
        phone.set("")
        email.set("")
        address.set("")

    def load_data():
        for item in table.get_children():
            table.delete(item)

        wb = load_workbook(FILE_NAME)
        ws = wb["Customers"]

        for row in ws.iter_rows(min_row=2, values_only=True):
            table.insert("", tk.END, values=row)

        wb.close()

    def add_customer():
        if customer_id.get() == "" or name.get() == "" or phone.get() == "":
            messagebox.showwarning(
                "Warning",
                "Customer ID, Name and Phone are required."
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb["Customers"]

        # Check duplicate ID
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[0]) == customer_id.get():
                messagebox.showerror(
                    "Error",
                    "Customer ID already exists."
                )
                wb.close()
                return

        ws.append([
            customer_id.get(),
            name.get(),
            phone.get(),
            email.get(),
            address.get()
        ])

        wb.save(FILE_NAME)
        wb.close()

        messagebox.showinfo(
            "Success",
            "Customer added successfully."
        )

        clear_fields()
        load_data()

    def select_customer(event):
        selected = table.focus()

        if selected:
            values = table.item(selected, "values")

            customer_id.set(values[0])
            name.set(values[1])
            phone.set(values[2])
            email.set(values[3])
            address.set(values[4])

    def update_customer():
        if customer_id.get() == "":
            messagebox.showwarning(
                "Warning",
                "Select a customer first."
            )
            return

        wb = load_workbook(FILE_NAME)
        ws = wb["Customers"]

        found = False

        for row in ws.iter_rows(min_row=2):
            if str(row[0].value) == customer_id.get():

                row[1].value = name.get()
                row[2].value = phone.get()
                row[3].value = email.get()
                row[4].value = address.get()

                found = True
                break

        if found:
            wb.save(FILE_NAME)
            messagebox.showinfo(
                "Success",
                "Customer updated successfully."
            )
        else:
            messagebox.showerror(
                "Error",
                "Customer not found."
            )

        wb.close()
        clear_fields()
        load_data()

    def delete_customer():
        if customer_id.get() == "":
            messagebox.showwarning(
                "Warning",
                "Select a customer first."
            )
            return

        answer = messagebox.askyesno(
            "Delete",
            "Are you sure you want to delete this customer?"
        )

        if answer:
            wb = load_workbook(FILE_NAME)
            ws = wb["Customers"]

            found = False

            for row in ws.iter_rows(min_row=2):
                if str(row[0].value) == customer_id.get():
                    ws.delete_rows(row[0].row, 1)
                    found = True
                    break

            if found:
                wb.save(FILE_NAME)
                messagebox.showinfo(
                    "Success",
                    "Customer deleted successfully."
                )
            else:
                messagebox.showerror(
                    "Error",
                    "Customer not found."
                )

            wb.close()
            clear_fields()
            load_data()

    def search_customer():
        search_text = search_var.get().lower()

        for item in table.get_children():
            table.delete(item)

        wb = load_workbook(FILE_NAME)
        ws = wb["Customers"]

        for row in ws.iter_rows(min_row=2, values_only=True):

            if any(
                search_text in str(value).lower()
                for value in row
            ):
                table.insert("", tk.END, values=row)

        wb.close()

    # -------- Heading --------

    tk.Label(
        root,
        text="CUSTOMER MANAGEMENT SYSTEM",
        font=("Arial", 22, "bold")
    ).pack(pady=15)

    # -------- Input Frame --------

    input_frame = tk.Frame(root)
    input_frame.pack(pady=5)

    tk.Label(input_frame, text="Customer ID").grid(
        row=0, column=0, padx=10, pady=8
    )
    tk.Entry(
        input_frame,
        textvariable=customer_id,
        width=25
    ).grid(row=0, column=1)

    tk.Label(input_frame, text="Name").grid(
        row=1, column=0, padx=10, pady=8
    )
    tk.Entry(
        input_frame,
        textvariable=name,
        width=25
    ).grid(row=1, column=1)

    tk.Label(input_frame, text="Phone").grid(
        row=2, column=0, padx=10, pady=8
    )
    tk.Entry(
        input_frame,
        textvariable=phone,
        width=25
    ).grid(row=2, column=1)

    tk.Label(input_frame, text="Email").grid(
        row=0, column=2, padx=10, pady=8
    )
    tk.Entry(
        input_frame,
        textvariable=email,
        width=25
    ).grid(row=0, column=3)

    tk.Label(input_frame, text="Address").grid(
        row=1, column=2, padx=10, pady=8
    )
    tk.Entry(
        input_frame,
        textvariable=address,
        width=25
    ).grid(row=1, column=3)

    # -------- Buttons --------

    button_frame = tk.Frame(root)
    button_frame.pack(pady=15)

    tk.Button(
        button_frame,
        text="Add Customer",
        width=15,
        command=add_customer
    ).grid(row=0, column=0, padx=5)

    tk.Button(
        button_frame,
        text="Update",
        width=15,
        command=update_customer
    ).grid(row=0, column=1, padx=5)

    tk.Button(
        button_frame,
        text="Delete",
        width=15,
        command=delete_customer
    ).grid(row=0, column=2, padx=5)

    tk.Button(
        button_frame,
        text="Reset",
        width=15,
        command=clear_fields
    ).grid(row=0, column=3, padx=5)

    # -------- Search --------

    search_frame = tk.Frame(root)
    search_frame.pack(pady=5)

    tk.Label(
        search_frame,
        text="Search:"
    ).pack(side=tk.LEFT, padx=5)

    tk.Entry(
        search_frame,
        textvariable=search_var,
        width=30
    ).pack(side=tk.LEFT, padx=5)

    tk.Button(
        search_frame,
        text="Search",
        command=search_customer
    ).pack(side=tk.LEFT)

    tk.Button(
        search_frame,
        text="Show All",
        command=load_data
    ).pack(side=tk.LEFT, padx=5)

    # -------- Customer Table --------

    table_frame = tk.Frame(root)
    table_frame.pack(fill=tk.BOTH, expand=True, padx=20, pady=10)

    columns = (
        "Customer ID",
        "Name",
        "Phone",
        "Email",
        "Address"
    )

    table = ttk.Treeview(
        table_frame,
        columns=columns,
        show="headings"
    )

    for column in columns:
        table.heading(column, text=column)
        table.column(column, width=150)

    table.pack(
        side=tk.LEFT,
        fill=tk.BOTH,
        expand=True
    )

    scrollbar = ttk.Scrollbar(
        table_frame,
        orient=tk.VERTICAL,
        command=table.yview
    )

    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    table.configure(
        yscrollcommand=scrollbar.set
    )

    table.bind(
        "<ButtonRelease-1>",
        select_customer
    )

    # Load existing customers
    load_data()

    root.mainloop()


# Create Excel file before starting
create_file()

login_window.mainloop()